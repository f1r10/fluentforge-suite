from pathlib import Path

import fitz
from docx import Document

from app.extractors import detect_candidates, extract_document


def test_plain_text_question_detection(tmp_path: Path):
    source = tmp_path / "questions.txt"
    source.write_text(
        "1. Capital of France?\n"
        "A. Paris\n"
        "B. London\n"
        "Answer: 1 A\n",
        encoding="utf-8",
    )

    extraction = extract_document(source, source.name, "text/plain")
    items = detect_candidates(extraction)

    assert extraction.method == "native_text"
    assert len(items) == 1
    assert items[0]["item_type"] == "question"
    assert items[0]["payload"]["prompt"] == "Capital of France?"
    assert items[0]["payload"]["answer_key"] == {"correct": ["a"]}


def test_docx_question_detection(tmp_path: Path):
    source = tmp_path / "questions.docx"
    document = Document()
    document.add_paragraph("1. Which option is correct?")
    document.add_paragraph("A. First")
    document.add_paragraph("B. Second")
    document.add_paragraph("Answer: 1 B")
    document.save(source)

    extraction = extract_document(
        source,
        source.name,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    items = detect_candidates(
        extraction,
        profile={"expected_content": "questions"},
    )

    assert extraction.method == "docx_native"
    assert len(items) == 1
    assert items[0]["item_type"] == "question"
    assert items[0]["payload"]["prompt"] == "Which option is correct?"
    assert items[0]["payload"]["answer_key"] == {"correct": ["b"]}


def test_csv_question_detection(tmp_path: Path):
    source = tmp_path / "questions.csv"
    source.write_text(
        "question,type,option_a,option_b,correct\n"
        "2+2?,single_choice,4,5,A\n",
        encoding="utf-8",
    )

    extraction = extract_document(source, source.name, "text/csv")
    items = detect_candidates(extraction)

    assert extraction.method == "delimited_native"
    assert len(items) == 1
    assert items[0]["payload"]["question_type"] == "single_choice"
    assert items[0]["payload"]["answer_key"] == {"correct": ["a"]}


def test_falls_back_to_raw_text_when_no_question_pattern(tmp_path: Path):
    source = tmp_path / "notes.txt"
    source.write_text("A paragraph without a numbered question.", encoding="utf-8")

    extraction = extract_document(source, source.name, "text/plain")
    items = detect_candidates(extraction)

    assert len(items) == 1
    assert items[0]["item_type"] == "raw_text"


def test_pdf_layout_blocks_include_normalized_crop(tmp_path: Path):
    source = tmp_path / "layout.pdf"
    document = fitz.open()
    page = document.new_page(width=600, height=800)
    page.insert_text((72, 100), "1. Example question")
    page.insert_text((72, 130), "A. First option")
    page.insert_text((72, 160), "B. Second option")
    document.save(source)
    document.close()

    extraction = extract_document(source, source.name, "application/pdf")

    assert extraction.method == "pdf_layout_native"
    assert extraction.stats["pages"] == 1
    assert extraction.pages[0]["blocks"]
    crop = extraction.pages[0]["blocks"][0]["crop"]
    assert 0 <= crop["x"] <= 1
    assert 0 <= crop["y"] <= 1
    assert 0 < crop["width"] <= 1
    assert 0 < crop["height"] <= 1


def test_pdf_reconstructs_reading_and_links_questions(tmp_path: Path):
    source = tmp_path / "reading.pdf"
    document = fitz.open()
    page = document.new_page(width=600, height=800)

    page.insert_text((72, 70), "A Short Passage")
    passage = (
        "Learning another language takes regular practice and repeated exposure. "
        "Students improve more quickly when they read meaningful texts, review "
        "new vocabulary, listen to authentic speech, and answer questions about "
        "what they understood. Teachers can then use those answers to identify "
        "weak areas and choose the next activity."
    )
    page.insert_textbox(fitz.Rect(72, 95, 530, 250), passage, fontsize=11)

    page.insert_text((72, 330), "1. What helps students improve?")
    page.insert_text((92, 360), "A. Regular practice")
    page.insert_text((92, 390), "B. Avoiding reading")
    page.insert_text((72, 440), "2. What can teachers identify?")
    page.insert_text((92, 470), "A. Weak areas")
    page.insert_text((92, 500), "B. Nothing")
    page.insert_text((72, 570), "Answers: 1 A 2 A")

    document.save(source)
    document.close()

    extraction = extract_document(source, source.name, "application/pdf")
    items = detect_candidates(extraction)

    readings = [item for item in items if item["item_type"] == "reading"]
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(readings) == 1
    assert len(questions) == 2
    assert readings[0]["payload"]["title"] == "A Short Passage"
    assert "Learning another language" in readings[0]["payload"]["body"]

    source_ref = readings[0]["payload"]["source_ref"]
    for index, question in enumerate(questions):
        context = question["payload"]["import_context"]
        assert context["kind"] == "reading"
        assert context["source_ref"] == source_ref
        assert context["sort_order"] == index
        assert question["crop"] is not None


def test_pdf_short_intro_does_not_create_false_reading(tmp_path: Path):
    source = tmp_path / "short-intro.pdf"
    document = fitz.open()
    page = document.new_page(width=600, height=800)
    page.insert_text((72, 80), "Choose the correct answer.")
    page.insert_text((72, 140), "1. Capital of France?")
    page.insert_text((92, 170), "A. Paris")
    page.insert_text((92, 200), "B. London")
    page.insert_text((72, 250), "Answer: 1 A")
    document.save(source)
    document.close()

    extraction = extract_document(source, source.name, "application/pdf")
    items = detect_candidates(extraction)

    assert not [item for item in items if item["item_type"] == "reading"]
    assert len([item for item in items if item["item_type"] == "question"]) == 1


def test_custom_spreadsheet_mapping_uses_header_row_fields_and_section(
    tmp_path: Path,
):
    source = tmp_path / "mapped.csv"
    source.write_text(
        "metadata,,,,,,,,,,,,,,\n"
        "Prompt Text,Kind,Choice One,Choice Two,Correct Value,Teacher Note,Rationale,Marks,Diff,Lang,CEFR,Labels,Unit\n"
        "Pick both,multiple_choice,Alpha,Beta,A|Beta,Choose all,Two valid choices,3.5,4,en,B2,grammar|exam,Unit 7\n",
        encoding="utf-8",
    )

    extraction = extract_document(source, source.name, "text/csv")
    profile = {
        "spreadsheet_mapping": {
            "header_row": 2,
            "first_data_row": 3,
            "include_sheets": ["mapped"],
            "sheet_as_section": False,
            "multi_value_separator": "|",
            "columns": {
                "prompt": "Prompt Text",
                "question_type": "Kind",
                "correct_answer": "Correct Value",
                "option_a": "Choice One",
                "option_b": "Choice Two",
                "instructions": "Teacher Note",
                "explanation": "Rationale",
                "points": "Marks",
                "difficulty": "Diff",
                "learning_language": "Lang",
                "level": "CEFR",
                "tags": "Labels",
                "section": "Unit",
            },
        }
    }

    items = detect_candidates(extraction, profile=profile)

    assert len(items) == 1
    item = items[0]
    assert item["item_type"] == "question"
    assert item["sheet"] == "mapped"
    assert item["payload"]["prompt"] == "Pick both"
    assert item["payload"]["answer_key"] == {"correct": ["a", "b"]}
    assert item["payload"]["instructions"] == "Choose all"
    assert item["payload"]["explanation"] == "Two valid choices"
    assert item["payload"]["scoring"]["points"] == 3.5
    assert item["payload"]["difficulty"] == 4
    assert item["payload"]["learning_language"] == "en"
    assert item["payload"]["level"] == "B2"
    assert item["payload"]["tags"] == ["grammar", "exam"]
    assert item["payload"]["import_context"]["section"] == "Unit 7"
    assert item["payload"]["import_mapping"]["row"] == 3


def test_custom_spreadsheet_mapping_respects_sheet_allowlist(tmp_path: Path):
    import openpyxl

    source = tmp_path / "multi.xlsx"
    workbook = openpyxl.Workbook()
    keep = workbook.active
    keep.title = "Keep"
    keep.append(["Q", "A"])
    keep.append(["Allowed question", "yes"])

    skip = workbook.create_sheet("Skip")
    skip.append(["Q", "A"])
    skip.append(["Excluded question", "no"])
    workbook.save(source)

    extraction = extract_document(
        source,
        source.name,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    profile = {
        "spreadsheet_mapping": {
            "header_row": 1,
            "first_data_row": 2,
            "include_sheets": ["Keep"],
            "sheet_as_section": True,
            "multi_value_separator": "|",
            "columns": {
                "prompt": "Q",
                "question_type": "",
                "correct_answer": "A",
                "option_a": "",
                "option_b": "",
                "option_c": "",
                "option_d": "",
                "option_e": "",
                "option_f": "",
                "option_g": "",
                "option_h": "",
                "instructions": "",
                "explanation": "",
                "points": "",
                "difficulty": "",
                "learning_language": "",
                "level": "",
                "tags": "",
                "section": "",
            },
        }
    }

    items = detect_candidates(extraction, profile=profile)

    assert len(items) == 1
    assert items[0]["sheet"] == "Keep"
    assert items[0]["payload"]["prompt"] == "Allowed question"
    assert items[0]["payload"]["answer_key"] == {"blanks": [["yes"]]}
    assert items[0]["payload"]["import_context"]["section"] == "Keep"


def test_custom_mapping_does_not_fallback_to_raw_text_when_selected_sheet_has_no_rows(
    tmp_path: Path,
):
    source = tmp_path / "questions.csv"
    source.write_text("Q,A\nQuestion,Answer\n", encoding="utf-8")
    extraction = extract_document(source, source.name, "text/csv")

    profile = {
        "spreadsheet_mapping": {
            "header_row": 1,
            "first_data_row": 2,
            "include_sheets": ["different-sheet"],
            "sheet_as_section": False,
            "multi_value_separator": "|",
            "columns": {"prompt": "Q"},
        }
    }

    items = detect_candidates(extraction, profile=profile)

    assert items == []


def test_inline_pdf_style_options_are_split():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    "SECTION I: Grammar\n"
                    "1. George is ................ than Nick.\n"
                    "a) tall b) taller c) tallest\n"
                    "2. What time ..... Calais tomorrow afternoon? "
                    "a) do the ferry reach b) is the ferry reaching c) does the ferry reach\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(questions) == 2
    assert [option["text"] for option in questions[0]["payload"]["payload"]["options"]] == [
        "tall",
        "taller",
        "tallest",
    ]
    assert questions[1]["payload"]["prompt"] == "What time ..... Calais tomorrow afternoon?"
    assert len(questions[1]["payload"]["payload"]["options"]) == 3


def test_explicit_listening_import_creates_context_and_links_questions():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 4,
                "text": (
                    "Listening Task 1\n"
                    "Listen to the recording and choose the correct answer.\n"
                    "1. Where is the speaker going?\n"
                    "A) London B) Paris C) Rome\n"
                    "2. When does the train leave?\n"
                    "A) Monday B) Tuesday C) Wednesday\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "listenings"})
    listenings = [item for item in items if item["item_type"] == "listening"]
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(listenings) == 1
    assert len(questions) == 2
    source_ref = listenings[0]["payload"]["source_ref"]
    assert listenings[0]["payload"]["metadata"]["audio_required"] is True
    for question in questions:
        assert question["payload"]["import_context"]["kind"] == "listening"
        assert question["payload"]["import_context"]["source_ref"] == source_ref


def test_questions_target_does_not_create_reading_context():
    from app.extractors import Extraction

    passage = (
        "A long reading passage with enough words to normally be detected as reading context. "
        "It contains several sentences so that the automatic reading heuristic would ordinarily "
        "create a reading entity before the questions that follow. Students should read it carefully."
    )
    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    f"{passage}\n"
                    "1. What should students do?\n"
                    "A) Read carefully B) Ignore it\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})

    assert not [item for item in items if item["item_type"] == "reading"]
    questions = [item for item in items if item["item_type"] == "question"]
    assert len(questions) == 1
    assert questions[0]["payload"].get("import_context", {}).get("kind") != "reading"


def test_checkbox_pdf_options_and_true_false_are_reconstructed():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    "1. Do you work on Saturdays? A \uf072 Yes, I work B \uf072 Yes, I do C \uf072 Yes, I am\n"
                    "2. The post office is in their house. ❏ True ❏ False\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(questions) == 2
    assert [option["text"] for option in questions[0]["payload"]["payload"]["options"]] == [
        "Yes, I work",
        "Yes, I do",
        "Yes, I am",
    ]
    assert questions[1]["payload"]["payload"]["options"] == [
        {"id": "true", "text": "True"},
        {"id": "false", "text": "False"},
    ]


def test_number_on_own_line_and_nested_numbered_statements_stay_one_question():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    "1.\n"
                    "Choose the correct variant.\n"
                    "How … to get to the airport?\n"
                    "1. long it is\n"
                    "2. did you\n"
                    "3. are you going\n"
                    "4. long does it take you\n"
                    "5. much does it\n"
                    "A) 1, 2\nB) 2, 3\nC) 3, 4\nD) 4, 5\nE) 1, 5\n"
                    "2.\n"
                    "Choose the correct variant.\n"
                    "This is my nephew.\n"
                    "A) brother B) cousin C) nephew\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(questions) == 2
    assert "1. long it is" in questions[0]["payload"]["prompt"]
    assert "5. much does it" in questions[0]["payload"]["prompt"]
    assert len(questions[0]["payload"]["payload"]["options"]) == 5
    assert questions[1]["payload"]["prompt"].startswith("Choose the correct variant.")


def test_answer_section_numbered_rows_are_not_imported_as_questions():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    "1. I ____ for Mike. (to work)\n"
                    "2. She ____ here. (to work)\n"
                    "Answers to the present simple test\n"
                    "1. work\n"
                    "2. works\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(questions) == 2
    assert questions[0]["payload"]["prompt"] == "I ____ for Mike. (to work)"
    assert questions[1]["payload"]["prompt"] == "She ____ here. (to work)"


def test_private_use_checkbox_glyph_is_normalized():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    "1. Do you work on Saturdays?\n"
                    "A \uf020 Yes, I work B \uf020 Yes, I do C \uf020 Yes, I am\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    question = [item for item in items if item["item_type"] == "question"][0]

    assert [option["text"] for option in question["payload"]["payload"]["options"]] == [
        "Yes, I work",
        "Yes, I do",
        "Yes, I am",
    ]


def test_inline_main_question_can_contain_numbered_statements():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 2,
                "text": (
                    "11. Choose the correct variant.\n"
                    "When Peter …, he … ever listened to.\n"
                    "1. speaks, has not\n"
                    "2. spoke, had not\n"
                    "3. speaks, is not\n"
                    "4. spoke, was not\n"
                    "A) 2, 3\nB) 1, 4\nC) 1, 2\nD) 2, 4\nE) 3, 4\n"
                    "12. Choose the correct tense form.\n"
                    "After Linda … a driving test, she … a car.\n"
                    "1. has passed, buy\n2. passed, bought\n"
                    "A) 1, 2\nB) 2, 1\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(questions) == 2
    assert "4. spoke, was not" in questions[0]["payload"]["prompt"]
    assert len(questions[0]["payload"]["payload"]["options"]) == 5
    assert "2. passed, bought" in questions[1]["payload"]["prompt"]


def test_wrapped_option_text_is_joined_and_section_heading_is_not_appended():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    "6. She's a doctor.\n"
                    "A \uf020 What's his job? B \uf020\n"
                    "What's your job?\n"
                    "C \uf020 What's her job?\n"
                    "Exercise 3: Prepositions.\n"
                    "7. ___ the summer, we go to the beach.\n"
                    "A) In B) At\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(questions) == 2
    assert [option["text"] for option in questions[0]["payload"]["payload"]["options"]] == [
        "What's his job?",
        "What's your job?",
        "What's her job?",
    ]
    assert "Exercise 3" not in questions[0]["payload"]["prompt"]


def test_global_answer_key_fills_unique_choice_and_short_answers():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    "1. Pick one.\nA) Alpha B) Beta C) Gamma\n"
                    "2. I ____ here. (to work)\n"
                ),
            },
            {
                "page": 2,
                "text": (
                    "Answer Key:\n"
                    "1: B\n"
                    "2. work\n"
                ),
            },
        ],
        [],
        (
            "1. Pick one.\nA) Alpha B) Beta C) Gamma\n"
            "2. I ____ here. (to work)\n"
            "Answer Key:\n1: B\n2. work\n"
        ),
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    questions = [item for item in items if item["item_type"] == "question"]

    assert len(questions) == 2
    assert questions[0]["payload"]["answer_key"] == {"correct": ["b"]}
    assert questions[1]["payload"]["answer_key"] == {"blanks": [["work"]]}


def test_standalone_true_false_lines_become_choices():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 1,
                "text": (
                    "1. The restaurant was expensive.\n"
                    "❏ True\n"
                    "❏ False\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "questions"})
    question = [item for item in items if item["item_type"] == "question"][0]

    assert question["payload"]["payload"]["options"] == [
        {"id": "true", "text": "True"},
        {"id": "false", "text": "False"},
    ]


def test_listening_target_splits_multiple_tasks_on_one_page():
    from app.extractors import Extraction

    extraction = Extraction(
        "native_text",
        [
            {
                "page": 9,
                "text": (
                    "SECTION III: Listening Comprehension\n"
                    "Listening Task 1\n"
                    "Listen and mark True or False.\n"
                    "1. Brazilians like coffee.\n❏ True\n❏ False\n"
                    "2. Filipinos eat rice.\n❏ True\n❏ False\n"
                    "Listening Task 2\n"
                    "Listen and mark True or False.\n"
                    "1. Snakes can hear.\n❏ True\n❏ False\n"
                    "2. Penguins can swim.\n❏ True\n❏ False\n"
                ),
            }
        ],
        [],
        "",
        {},
    )

    items = detect_candidates(extraction, profile={"expected_content": "listenings"})
    listenings = [item for item in items if item["item_type"] == "listening"]
    questions = [item for item in items if item["item_type"] == "question"]

    assert [item["payload"]["title"] for item in listenings] == [
        "Listening Task 1",
        "Listening Task 2",
    ]
    assert len(questions) == 4
    assert {
        question["payload"]["import_context"]["source_ref"]
        for question in questions
    } == {
        "listening:page:9:task:1",
        "listening:page:9:task:2",
    }
